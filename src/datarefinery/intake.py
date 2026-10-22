"""Bounded intake: raw bytes to a string-typed frame plus an immutable source manifest. No coercion happens here."""

from __future__ import annotations

import csv
import hashlib
import io
import json
from pathlib import Path

import openpyxl
import pandas as pd

MAX_ROWS = 5000
MAX_BYTES = 2_000_000
MAX_SHEETS = 1
FORMATS = {".csv": "csv", ".json": "json", ".xlsx": "xlsx"}


class IntakeError(ValueError):
    pass


def find_source(fixtures_dir: Path, batch: str) -> Path:
    found = sorted(p for p in fixtures_dir.glob(f"orders_{batch}.*") if p.suffix in FORMATS)
    if len(found) != 1:
        raise IntakeError(f"batch {batch}: expected exactly one orders_{batch}.(csv|json|xlsx), found {len(found)}")
    return found[0]


def _text(value: object) -> str | None:
    if value is None or isinstance(value, str):
        return value
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float):
        return repr(value)
    return str(value)


def _csv(data: bytes) -> tuple[list[str], list[list[str | None]]]:
    rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"), newline="")))
    if not rows:
        raise IntakeError("empty CSV")
    header, body = rows[0], rows[1:]
    for n, row in enumerate(body, start=1):
        if len(row) != len(header):
            raise IntakeError(f"CSV record {n} has {len(row)} fields, header has {len(header)}")
    return header, body


def _json(data: bytes) -> tuple[list[str], list[list[str | None]]]:
    # Numbers stay as their literal text so 12.50 is not silently turned into 12.5.
    records = json.loads(data.decode("utf-8"), parse_float=str, parse_int=str)
    if not isinstance(records, list) or not all(isinstance(r, dict) for r in records):
        raise IntakeError("JSON batch must be an array of objects")
    header: list[str] = []
    for record in records:
        header += [k for k in record if k not in header]
    body = []
    for record in records:
        if any(isinstance(v, dict | list) for v in record.values()):
            raise IntakeError("JSON batch values must be scalars")
        body.append([_text(record.get(k)) for k in header])
    return header, body


def _xlsx(data: bytes) -> tuple[list[str], list[list[str | None]]]:
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        if len(wb.sheetnames) > MAX_SHEETS:
            raise IntakeError(f"XLSX has {len(wb.sheetnames)} sheets, max {MAX_SHEETS}")
        rows = [[_text(v) for v in row] for row in wb[wb.sheetnames[0]].iter_rows(values_only=True)]
    finally:
        wb.close()
    if not rows:
        raise IntakeError("empty XLSX sheet")
    return [str(h) for h in rows[0]], rows[1:]


def read_batch(path: Path) -> tuple[pd.DataFrame, dict]:
    """Frame indexed by 1-based source record number (header excluded), every cell str or None."""
    fmt = FORMATS.get(path.suffix)
    if fmt is None:
        raise IntakeError(f"{path.name}: unsupported format")
    data = path.read_bytes()
    if len(data) > MAX_BYTES:
        raise IntakeError(f"{path.name}: {len(data)} bytes exceeds the {MAX_BYTES} byte bound")
    header, body = {"csv": _csv, "json": _json, "xlsx": _xlsx}[fmt](data)
    if len(set(header)) != len(header) or any(not h for h in header):
        raise IntakeError(f"{path.name}: header has blank or repeated column names")
    if len(body) > MAX_ROWS:
        raise IntakeError(f"{path.name}: {len(body)} rows exceeds the {MAX_ROWS} row bound")
    frame = pd.DataFrame(body, columns=header, dtype=object, index=pd.RangeIndex(1, len(body) + 1))
    manifest = {
        "file": path.name,
        "format": fmt,
        "sha256": hashlib.sha256(data).hexdigest(),
        "bytes": len(data),
        "rows": len(body),
        "columns": header,
    }
    return frame, manifest
